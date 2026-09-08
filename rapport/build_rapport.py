#!/usr/bin/env python3
"""Construit rapport-pfa.docx avec une mise en page professionnelle.

Chaîne : rapport-pfa.md  --(pandoc + modèle de styles)-->  corps .docx
         puis ajout programmatique de la page de garde, de la table des
         matières et de la mise en forme des tableaux (python-docx).

Le modèle de styles (couleurs de titres, polices Calibri/Cambria, marges,
pied de page numéroté) provient de rapport/assets/reference-style.docx.

Usage :
    .venv/bin/python rapport/build_rapport.py
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import pypdf
from docx import Document
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Cm, Twips

HERE = Path(__file__).resolve().parent
MD = HERE / "rapport-pfa.md"
REF = HERE / "assets" / "reference-style.docx"
LOGO = HERE / "assets" / "isga-logo.png"
BODY = HERE / ".body.docx"          # intermédiaire pandoc
OUT = HERE / "rapport-pfa.docx"

# ── Palette (thème Word « Office », identique au rapport de référence) ──
NAVY = RGBColor(0x17, 0x36, 0x5D)      # Titre
BLUE = RGBColor(0x36, 0x5F, 0x91)      # accent
GREY = RGBColor(0x59, 0x59, 0x59)      # libellés
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HDR_FILL = "1F3864"                     # entête de tableau (navy)
BAND_FILL = "DCE6F1"                    # lignes paires (bleu clair)

META = [
    ("Établissement", "ISGA — Institut Supérieur d'Ingénierie et des Affaires"),
    ("Filière / Niveau", "Cycle d'ingénieur — 2ᵉ année"),
    ("Réalisé par", "BENDADI Mohamed  ·  SADEK Zakaria  ·  BELHASSAN Amine"),
    ("Encadré par", "M. Adama SAMAKE"),
    ("Année universitaire", "2025 / 2026"),
]


def run_pandoc() -> None:
    subprocess.run(
        [
            "pandoc", str(MD),
            "--reference-doc", str(REF),
            "--resource-path", str(HERE),
            "-o", str(BODY),
        ],
        check=True,
        cwd=HERE,
    )


def set_cell_shading(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def clear_table_borders(tbl) -> None:
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        borders.append(el)
    tblPr.append(borders)


def set_repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))


def rebuild_table(doc, old_tbl, usable_twips: int):
    """Reconstruit un tableau propre (pandoc produit une grille incorrecte).

    Les tableaux issus de pandoc portent un style « Table » absent du modèle et
    une grille que LibreOffice/Word empilent. On recrée un tableau natif —
    entête navy (blanc gras), lignes zébrées, sans bordure — à l'image du
    rapport de référence, puis on remplace l'ancien.
    """
    from docx.enum.table import WD_TABLE_ALIGNMENT

    data = [[cell.text for cell in row.cells] for row in old_tbl.rows]
    nrows, ncols = len(data), len(data[0])

    new = doc.add_table(rows=nrows, cols=ncols)   # ajouté en fin de corps
    new.style = doc.styles["Table Grid"]
    new.alignment = WD_TABLE_ALIGNMENT.CENTER
    new.autofit = False

    tblPr = new._tbl.tblPr
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "dxa"); tblW.set(qn("w:w"), str(usable_twips))
    tblPr.append(tblW)
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    borders = OxmlElement("w:tblBorders")         # aucune bordure (comme la réf.)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)

    if ncols == 1:
        widths = [usable_twips]
    else:
        first = int(usable_twips * 0.42)
        rest = (usable_twips - first) // (ncols - 1)
        widths = [first] + [rest] * (ncols - 1)
    grid = new._tbl.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))

    for r in range(nrows):
        for c in range(ncols):
            cell = new.rows[r].cells[c]
            cell.width = Twips(widths[c])
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.space_after = Pt(3)
            run = para.add_run(data[r][c])
            run.font.name = "Calibri"; run.font.size = Pt(10)
            if r == 0:
                run.font.bold = True; run.font.color.rgb = WHITE
                set_cell_shading(cell, HDR_FILL)
            elif r % 2 == 1:
                set_cell_shading(cell, BAND_FILL)
    set_repeat_header(new.rows[0])

    old_tbl._tbl.addprevious(new._tbl)            # placer avant l'ancien...
    old_tbl._tbl.getparent().remove(old_tbl._tbl)  # ...puis le supprimer
    return new


def style_code_blocks(doc) -> None:
    """Bloc de code : police à chasse fixe + fond gris clair (comme la réf.)."""
    for p in doc.paragraphs:
        if not (p.style and p.style.name == "Source Code"):
            continue
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F3F5")
        pPr.append(shd)
        for run in p.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(9)


def style_figure_captions(doc) -> None:
    """Légendes « Figure … » en style Caption (bleu, centré) + image centrée."""
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        t = p.text.strip()
        if t.startswith("Figure") and t[6:7] in " 0123456789":
            try:
                p.style = doc.styles["Caption"]
            except KeyError:
                pass
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i > 0:                       # centrer l'image qui précède
                paras[i - 1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_page_break_before(anchor) -> None:
    p = anchor.insert_paragraph_before()
    p.add_run().add_break(WD_BREAK.PAGE)


def build_cover(doc, anchor) -> None:
    """Insère la page de garde juste avant `anchor` (1er élément du corps)."""
    # 1) logo centré
    p = anchor.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(72)
    p.paragraph_format.space_after = Pt(28)
    p.add_run().add_picture(str(LOGO), width=Cm(6.2))

    # 2) titre principal (style Title, centré)
    p = anchor.insert_paragraph_before(
        "Système intelligent de gestion des présences\npar reconnaissance faciale"
    )
    try:
        p.style = doc.styles["Title"]
    except KeyError:
        pass
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 3) ligne d'accent
    p = anchor.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("« AI Class Attendance for Students »")
    r.font.size = Pt(17); r.font.color.rgb = BLUE; r.font.name = "Calibri"

    # 4) filet horizontal
    rule = anchor.insert_paragraph_before()
    rule.paragraph_format.space_before = Pt(10)
    pPr = rule._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "365F91")
    pbdr.append(bottom); pPr.append(pbdr)

    # 5) sous-titre italique bleu (style Subtitle)
    p = anchor.insert_paragraph_before(
        "Rapport de Projet de Fin d'Année"
    )
    try:
        p.style = doc.styles["Subtitle"]
    except KeyError:
        pass
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 6) bloc d'informations (tableau 2 colonnes sans bordure)
    tbl = doc.add_table(rows=len(META), cols=2)
    clear_table_borders(tbl)
    for i, (label, value) in enumerate(META):
        c0, c1 = tbl.rows[i].cells
        c0.width = Cm(4.6); c1.width = Cm(10.4)
        pl = c0.paragraphs[0]; pl.paragraph_format.space_after = Pt(6)
        rl = pl.add_run(label)
        rl.font.size = Pt(10.5); rl.font.color.rgb = GREY; rl.font.name = "Calibri"
        pv = c1.paragraphs[0]; pv.paragraph_format.space_after = Pt(6)
        rv = pv.add_run(value)
        rv.font.size = Pt(10.5); rv.font.bold = True; rv.font.name = "Calibri"
    tbl_holder = anchor.insert_paragraph_before()  # espace avant tableau
    tbl_holder.paragraph_format.space_before = Pt(60)
    anchor._p.addprevious(tbl._tbl)

    # 7) saut de page
    add_page_break_before(anchor)


_LEVEL = {"heading 1": 1, "heading 2": 2, "heading 3": 3}


def collect_toc_entries(doc, start_anchor):
    """(niveau, texte) des titres H1–H3 à partir de `start_anchor` (chapitres)."""
    entries, started = [], False
    for p in doc.paragraphs:
        if p._p is start_anchor._p:      # comparer l'élément XML (les wrappers diffèrent)
            started = True
        if not started or not p.style:
            continue
        lvl = _LEVEL.get(p.style.name.lower())
        if lvl and p.text.strip():
            entries.append((lvl, p.text.strip()))
    return entries


def build_toc(doc, anchor) -> None:
    """Insère « Table des matières » + entrées à points de conduite avant `anchor`."""
    sec = doc.sections[0]
    tab_pos = sec.page_width - sec.left_margin - sec.right_margin
    entries = collect_toc_entries(doc, anchor)

    add_page_break_before(anchor)
    head = anchor.insert_paragraph_before("Table des matières")
    try:
        head.style = doc.styles["TOC Heading"]
    except KeyError:
        head.style = doc.styles["Heading 1"]

    for lvl, text in entries:
        p = anchor.insert_paragraph_before()
        try:
            p.style = doc.styles[f"toc {lvl}"]
        except KeyError:
            pass
        p.paragraph_format.tab_stops.add_tab_stop(
            tab_pos, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        p.add_run(text)
        p.add_run("\t")        # rempli plus tard par le numéro de page
    add_page_break_before(anchor)


def bake_toc_pages(out_path: Path) -> None:
    """Rend le document, lit le sommaire PDF et inscrit les numéros de page."""
    if not shutil.which("soffice"):
        print("! soffice introuvable : numéros de page du sommaire non calculés")
        return
    tmp = Path(tempfile.mkdtemp())
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp),
         str(out_path)],
        check=True, capture_output=True,
    )
    pdf = tmp / (out_path.stem + ".pdf")
    reader = pypdf.PdfReader(str(pdf))

    outline = []

    def walk(items):
        for it in items:
            if isinstance(it, list):
                walk(it)
            else:
                outline.append((it.title.strip(),
                                reader.get_destination_page_number(it) + 1))
    walk(reader.outline)

    doc = Document(str(out_path))
    toc_paras = [p for p in doc.paragraphs
                 if p.style and p.style.name.lower() in ("toc 1", "toc 2", "toc 3")]
    # aligner : le sommaire PDF liste aussi le liminaire ; on cale sur le 1er titre TOC
    first = toc_paras[0].runs[0].text.strip()
    start = next((i for i, (t, _) in enumerate(outline) if t == first), 0)
    pages = [pg for _, pg in outline[start:]]

    if len(pages) != len(toc_paras):
        print(f"! sommaire : {len(toc_paras)} entrées vs {len(pages)} pages PDF")
    for para, page in zip(toc_paras, pages):
        para.runs[-1].text = "\t" + str(page)
    doc.save(str(out_path))
    shutil.rmtree(tmp, ignore_errors=True)
    print(f"sommaire : {len(pages)} entrées paginées")


def first_chapter_paragraph(doc):
    for p in doc.paragraphs:
        if p.style and p.style.name.lower().startswith("heading 1") \
                and p.text.strip().lower().startswith("chapitre"):
            return p
    raise RuntimeError("Chapitre 1 introuvable")


def main() -> None:
    run_pandoc()
    doc = Document(str(BODY))

    sec = doc.sections[0]
    usable = int((sec.page_width - sec.left_margin - sec.right_margin) / 635)  # EMU -> twips
    for tbl in list(doc.tables):      # snapshot : on remplace chaque tableau du contenu
        rebuild_table(doc, tbl, usable)
    style_figure_captions(doc)
    style_code_blocks(doc)

    chap1 = first_chapter_paragraph(doc)
    build_toc(doc, chap1)

    body_first = doc.paragraphs[0]    # « Remerciements »
    build_cover(doc, body_first)

    doc.save(str(OUT))
    BODY.unlink(missing_ok=True)
    bake_toc_pages(OUT)               # 2ᵉ passe : numéros de page du sommaire
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    main()
