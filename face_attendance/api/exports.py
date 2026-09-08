"""Génération des rapports d'assiduité au format Excel et PDF."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

_INDIGO = "4F46E5"
_HEADERS = ["Étudiant", "Présences", "Séances", "Taux (%)"]


def report_xlsx(module_libelle: str, rows: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Assiduité"

    ws["A1"] = f"Rapport d'assiduité — {module_libelle}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(_HEADERS)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=_INDIGO)
        c.alignment = Alignment(horizontal="center")

    for r in rows:
        ws.append([r["nom"], r["presences"], r["total"], r["taux"]])

    for col, width in zip("ABCD", (32, 12, 10, 10)):
        ws.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_pdf(module_libelle: str, rows: list) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=f"Rapport — {module_libelle}")
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(f"Rapport d'assiduité — {module_libelle}", styles["Title"]),
        Spacer(1, 14),
    ]

    data = [_HEADERS] + [
        [r["nom"], str(r["presences"]), str(r["total"]), f'{r["taux"]}'] for r in rows
    ]
    table = Table(data, colWidths=[230, 80, 70, 70])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + _INDIGO)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                 [colors.white, colors.HexColor("#F4F5FB")]),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    return buf.getvalue()
